import csv,sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
S,O=sys.argv[1],sys.argv[2]
wb=Workbook(); first=True
def sheet(name,path,widths,fname_map=None):
    global first
    ws=wb.active if first else wb.create_sheet(); first=False; ws.title=name
    with open(path) as f:
        for i,row in enumerate(csv.reader(f,delimiter="\t")): ws.append(row)
    for c in ws[1]: c.font=Font(bold=True); c.fill=PatternFill("solid",fgColor="DDDDDD")
    for i,w in enumerate(widths): ws.column_dimensions[chr(65+i)].width=w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(wrap_text=True,vertical="top")
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
sheet("まとめ（同じ訳を1行）",f"{S}/ja-fixes-grouped.tsv",[30,26,30,8,6,40,10,6,50])
sheet("全件（キーごと）",f"{S}/ja-fixes-all.tsv",[22,40,30,26,30,8,6,40,10])
with open(f"{S}/missing.tsv","w") as f:
    f.write("ファイル\tキー\t英語\n")
    m={"common_locales":"common/locales/ja.yml","common_locales_enums":"common/locales/enums/ja.yml","frontend_config_locales":"frontend/config/locales/ja.yml","public_config_locales":"public/config/locales/ja.yml"}
    for r in csv.reader(open(f"{S}/out/missing.tsv"),delimiter="\t"): f.write(f"{m[r[0]]}\t{r[1]}\t{r[2]}\n")
sheet("未翻訳（キー自体が無い）",f"{S}/missing.tsv",[28,50,60])
wb.save(f"{O}/archivesspace-ja-fixes.xlsx")
